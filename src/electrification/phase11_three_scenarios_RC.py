import sys, io; sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import os as _os
sys.path.insert(0, _os.path.dirname(_os.path.dirname(_os.path.dirname(_os.path.abspath(__file__)))))
import config  # repository paths, see config.py
"""Phase 11 (Ref-Coupled): Three-scenario comparison with Ref-Elec coupling.

Same as phase11 but Ref1-5 cost/CI are auto-rescaled when P_elec or CI_elec
changes (i.e., S2 renewable PPA scenario also reduces Ref refrigeration cost
and CI proportionally).

Scenarios:
  S0: Baseline (P_elec=0.08, CI_elec=369)  → Ref at baseline
  S1: Grid electrification  (P_elec=0.08, CI_elec=369)  → Ref unchanged (same elec)
  S2: Renewable electrification (P_elec=0.04, CI_elec=30) → Ref scaled DOWN

Outputs (suffix _RC):
  - phase11_three_scenarios_RC.csv
  - fig_phase11_zone_summary_RC.png
"""
import sys, io, os, json
os.environ['PYTHONUTF8'] = '1'
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import load_workbook

# ─── Ref-Elec coupling helper ────────────────────────────────────────
from ref_elec_coupling import (
    ref_price, ref_ci,
    P_ELEC_BASE, CI_ELEC_BASE
)

BASE = config.INTERMEDIATE

# ════════════════════════════════════════════════════════════════════
# Load Global utility prices/CI (baseline) — for non-Ref carriers only
# ════════════════════════════════════════════════════════════════════
country_fp = config.COUNTRY_PRICES
wb = load_workbook(country_fp, data_only=True, read_only=True)
ws = wb['Sheet1']
PRICE_ROWS = {4: 'Elec', 5: 'FH', 6: 'CW',
              12: 'LP', 13: 'MP', 14: 'HP'}
CI_ROWS    = {17: 'Elec', 18: 'FH', 19: 'CW',
              25: 'LP', 26: 'MP', 27: 'HP'}
prices_base = {ws.cell(row=r, column=6).value: 0 for r in PRICE_ROWS}
ci_base     = {ws.cell(row=r, column=6).value: 0 for r in CI_ROWS}
prices_base = {}; ci_base = {}
for r, u in PRICE_ROWS.items(): prices_base[u] = ws.cell(row=r, column=6).value
for r, u in CI_ROWS.items():    ci_base[u]     = ws.cell(row=r, column=6).value
wb.close()

# Override: confirmed Elec CI
ci_base['Elec'] = 0.369 / 3.6   # kgCO2/MJ
print(f'(Override) Elec CI baseline = {ci_base["Elec"]:.4f} kgCO2/MJ = 0.369 kgCO2/kWh')

H2_BASELINE_PRICE = 2.0
H2_BASELINE_CI    = 7.0

# ════════════════════════════════════════════════════════════════════
# Scenario parameters
# ════════════════════════════════════════════════════════════════════
# Each scenario defines its electricity (used for both: carrier electrification
# AND Ref refrigeration coupling)
SCENARIOS = {
    'S0_Baseline': {
        'P_elec_kwh':  0.08,        # baseline
        'CI_elec_kwh': 0.369,       # baseline
        'electrify_carriers': False,
    },
    'S1_Grid_Electrified': {
        'P_elec_kwh':  0.08,        # grid (same as baseline)
        'CI_elec_kwh': 0.369,
        'electrify_carriers': True,
    },
    'S2_Renewable_Electrified': {
        'P_elec_kwh':  0.04,        # renewable PPA
        'CI_elec_kwh': 0.030,
        'electrify_carriers': True,
    },
}

E_H2       = 53.0
H2_CAPOP   = 0.7
ETA_BOILER = 0.99
ETA_HEATER = 0.95

def electrified_set(p_kwh, ci_kwh):
    """Per-MJ electrified carrier prices/CI from elec inputs (P_kWh, CI_kgCO2/kWh)."""
    p_mj  = p_kwh / 3.6
    ci_mj = ci_kwh / 3.6
    return {
        'H2_price':  E_H2 * p_kwh + H2_CAPOP,
        'H2_CI':     E_H2 * ci_kwh,
        'FH_price':  p_mj / ETA_HEATER,
        'FH_CI':     ci_mj / ETA_HEATER,
        'LP_price':  p_mj / ETA_BOILER,
        'LP_CI':     ci_mj / ETA_BOILER,
        'MP_price':  p_mj / ETA_BOILER,
        'MP_CI':     ci_mj / ETA_BOILER,
        'HP_price':  p_mj / ETA_BOILER,
        'HP_CI':     ci_mj / ETA_BOILER,
        'Elec_price': p_mj,
        'Elec_CI':    ci_mj,
    }

print('\n=== Scenario parameters ===')
for s, p in SCENARIOS.items():
    print(f'  {s}: P_elec={p["P_elec_kwh"]:.3f} $/kWh, CI={p["CI_elec_kwh"]*1000:.0f} g/kWh, '
          f'electrify={p["electrify_carriers"]}')

# ════════════════════════════════════════════════════════════════════
# Load 959 pathways
# ════════════════════════════════════════════════════════════════════
print('\n=== Loading All pathways ===')
df = pd.read_excel(config.MASTER_XLSX,
                   sheet_name='All pathways', header=3)
df['__feed'] = df.iloc[:, 0].astype(str).str.strip()
df['__tech'] = df.iloc[:, 1].astype(str).str.strip()
df = df[df['__feed'].isin(['PE','PET','PP','PS','PVC'])].copy().reset_index(drop=True)
print(f'Pathways: {len(df)}')

COST_COLS = {'Elec': 7, 'FH': 8, 'CW': 9,
             'REF1':10,'REF2':11,'REF3':12,'REF4':13,'REF5':14,
             'LP': 15, 'HP': 16, 'H2': 20}
GHG_COLS  = {'Elec':31, 'FH':32,
             'REF1':33,'REF2':34,'REF3':35,'REF4':36,'REF5':37,
             'LP':38,  'HP':39, 'H2':42}
MASS_TOTAL = 50e6

up_col = next(c for c in df.columns if isinstance(c, str) and c.lower().startswith('unit profit'))
cr_col = next(c for c in df.columns if isinstance(c, str) and c.lower().startswith('unit co2 reduc'))
df['__up_base'] = pd.to_numeric(df[up_col], errors='coerce')
df['__cr_base'] = pd.to_numeric(df[cr_col], errors='coerce')
df = df.dropna(subset=['__up_base', '__cr_base']).copy().reset_index(drop=True)
print(f'After dropna: {len(df)}')


def compute_scenario(name, params):
    """Compute UP and CR for given scenario with Ref-Elec coupling."""
    delta_cost = pd.Series(0.0, index=df.index)
    delta_ghg  = pd.Series(0.0, index=df.index)

    P_elec_kwh  = params['P_elec_kwh']
    CI_elec_kwh = params['CI_elec_kwh']

    # ─── Ref-Elec coupling (always applied, regardless of carrier electrification) ───
    # Ref refrigerators use the electricity at (P_elec_kwh, CI_elec_kwh).
    for k in range(1, 6):
        util = f'REF{k}'
        # Cost
        c_base = pd.to_numeric(df.iloc[:, COST_COLS[util]], errors='coerce').fillna(0)
        ratio_cost = ref_price(k, P_elec_kwh) / ref_price(k, P_ELEC_BASE)
        delta_cost += c_base * (ratio_cost - 1.0)
        # CI
        g_base = pd.to_numeric(df.iloc[:, GHG_COLS[util]], errors='coerce').fillna(0)
        ratio_ci = ref_ci(k, CI_elec_kwh*1000) / ref_ci(k, CI_ELEC_BASE)   # CI_elec in g/kWh
        delta_ghg += g_base * (ratio_ci - 1.0)

    # ─── Carrier electrification (H2, FH, LP, MP, HP, Elec itself) ───
    if params['electrify_carriers']:
        E = electrified_set(P_elec_kwh, CI_elec_kwh)

        for util, col in COST_COLS.items():
            if util.startswith('REF') or util == 'CW': continue   # Ref handled above; CW unchanged
            c_base_col = pd.to_numeric(df.iloc[:, col], errors='coerce').fillna(0)
            if util == 'H2':
                p_b, p_n = H2_BASELINE_PRICE, E['H2_price']
            else:
                p_b, p_n = prices_base[util], E[f'{util}_price']
            if p_b and p_b > 0:
                delta_cost += c_base_col * (p_n / p_b - 1.0)

        for util, col in GHG_COLS.items():
            if util.startswith('REF'): continue
            g_base_col = pd.to_numeric(df.iloc[:, col], errors='coerce').fillna(0)
            if util == 'H2':
                ci_b, ci_n = H2_BASELINE_CI, E['H2_CI']
            else:
                ci_b, ci_n = ci_base[util], E[f'{util}_CI']
            if ci_b and ci_b > 0:
                delta_ghg += g_base_col * (ci_n / ci_b - 1.0)

    new_up = df['__up_base'] - delta_cost * 1e6 / MASS_TOTAL
    new_cr = df['__cr_base'] - delta_ghg / MASS_TOTAL

    return pd.DataFrame({
        'idx':  np.arange(len(df)),
        'feed': df['__feed'],
        'tech': df['__tech'],
        'scenario': name,
        'UP':   new_up.values,
        'CR':   new_cr.values,
    })


print('\n=== Computing 3 scenarios (with Ref-Elec coupling) ===')
results = {name: compute_scenario(name, p) for name, p in SCENARIOS.items()}
all_three = pd.concat(results.values(), ignore_index=True)

def assign(p, c_):
    if p > -0.42 and c_ > 0: return 1
    if p > -0.42: return 2
    if c_ > 0: return 3
    return 4

all_three['zone'] = all_three.apply(lambda r: assign(r['UP'], r['CR']), axis=1)

print('\nZONE counts per scenario:')
print(all_three.groupby('scenario')['zone'].value_counts().unstack(fill_value=0))
print('\nUP mean/min/max per scenario:')
print(all_three.groupby('scenario')['UP'].describe()[['mean','min','max']].round(3))
print('\nCR mean/min/max per scenario:')
print(all_three.groupby('scenario')['CR'].describe()[['mean','min','max']].round(3))

# Save
out_csv = os.path.join(BASE, 'phase11_three_scenarios_RC.csv')
all_three.to_csv(out_csv, index=False, encoding='utf-8-sig')
print(f'\nSaved: {out_csv}')

# ── Zone stacked bar ─────────────────────────────────────────────
fig, ax = plt.subplots(figsize=(8, 6))
ZONE_COLORS = ['#27ae60', '#f39c12', '#3498db', '#c0392b']
ZONE_NAMES  = ['Z1 Win-Win', 'Z2 Econ', 'Z3 Env', 'Z4 Lose-Lose']
order = list(SCENARIOS.keys())
xlabels = ['S0\nBaseline', 'S1\nGrid', 'S2\nRenewable']
summary = (all_three.groupby('scenario')['zone'].value_counts()
                    .unstack(fill_value=0)
                    .reindex(index=order, columns=[1,2,3,4], fill_value=0))

x = np.arange(len(order))
bottom = np.zeros(len(order))
for j, (z, color, name) in enumerate(zip([1,2,3,4], ZONE_COLORS, ZONE_NAMES)):
    vals = summary[z].values
    ax.bar(x, vals, bottom=bottom, color=color, edgecolor='white', linewidth=0.8, label=name)
    for xi, v in enumerate(vals):
        if v > 25:
            ax.text(xi, bottom[xi] + v/2, f'{int(v)}', ha='center', va='center',
                    fontsize=10, color='white', fontweight='bold')
    bottom += vals

ax.set_xticks(x); ax.set_xticklabels(xlabels, fontsize=11)
ax.set_ylabel('Pathway count', fontsize=11)
ax.set_title('ZONE distribution (Ref-Elec coupled)', fontsize=12, pad=10)
ax.legend(loc='upper right', fontsize=9)
ax.grid(axis='y', linestyle=':', alpha=0.4)
plt.tight_layout()
out_z = os.path.join(config.FIG_OUT, 'fig_phase11_zone_summary_RC.png')
plt.savefig(out_z, dpi=200, bbox_inches='tight')
plt.close()
print(f'Saved: {out_z}')

print('\nDone.')
